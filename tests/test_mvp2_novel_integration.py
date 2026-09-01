"""Offline import-to-replay acceptance gate for faithful novel mode."""

from __future__ import annotations

import asyncio
from contextlib import redirect_stderr, redirect_stdout, suppress
from copy import deepcopy
import hashlib
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import warnings

from docx import Document
import httpx

import backend.main as main_module
from backend.provenance import dlc_provenance
from backend.story import (
    AnalysisStore,
    NovelSessionController,
    StorySourceLoader,
    offline_analysis_key,
)
from backend.story.import_analysis import main as import_analysis_main
from backend.story.planner import ChapterPlanner
from backend.story.source_store import PinnedStorySourceStore
from backend.story.story_map_codec import encode_story_map
from backend.timeline.models import CycleGapPolicy, SCHEMA_VERSION, SessionStatus
from tests.test_game_loop_timeline import make_game_loop_for_test
from tests.test_session_endpoints import make_endpoint_state


class LocalChapterPlannerClient:
    """The only model seam: it produces a selected-chapter plan in memory."""

    model = "mvp2-local-planner"

    def __init__(self) -> None:
        self.calls: list[dict[str, str]] = []

    async def complete_json(
        self, _system_prompt: str, user_content: str, schema_name: str
    ) -> dict[str, object]:
        self.calls.append(
            {
                "schema_name": schema_name,
                "system_prompt": _system_prompt,
                "user_content": user_content,
            }
        )
        request = json.loads(user_content)
        return {
            "scenes": [
                {
                    "scene_id": scene["scene_id"],
                    "channels": {
                        "A": {"mode": "set", "base_strength": 20},
                        "B": {"mode": "set", "base_strength": 10},
                    },
                }
                for scene in request["scenes"]
            ]
        }


class MVP2OfflineNovelIntegrationTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self) -> None:
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.harness = make_game_loop_for_test(
            self.root / "replays", autopilot_interval=3600, gap_tenths=(0, 0, 0)
        )
        self.cfg = deepcopy(self.harness.cfg)
        self.cfg["story"].update(
            {
                "import_dir": "data/stories",
                "analysis_dir": "data/story_analysis",
                "candidate_dir": "data/story_candidates",
            }
        )
        self.project_root_patch = patch.object(main_module, "PROJECT_ROOT", self.root)
        self.project_root_patch.start()
        self.addCleanup(self.project_root_patch.stop)

        self.state = make_endpoint_state(self.harness)
        self.state.cfg = self.cfg
        self.state.story_import_directory = self.root / self.cfg["story"]["import_dir"]
        self.state.story_source_store = PinnedStorySourceStore(
            self.state.story_import_directory, project_root=self.root
        )
        self.addCleanup(self.state.story_source_store.close)
        self.state.story_source_loader = StorySourceLoader(max_bytes=1024 * 1024)
        self.state.story_source_max_bytes = 1024 * 1024
        self.state.story_analysis_store = AnalysisStore(
            self.root / self.cfg["story"]["analysis_dir"]
        )
        self.state.story_sources = {}
        self.state.active_story_source_id = None
        self.state.story_source_generation = 0
        self.state.story_planning_task = None
        self.state.story_planning_context = None
        self.state.story_runtime_owner = None
        self.state.story_source_store_task = None
        self.state.story_source_io_tasks = set()
        self.state.story_cleanup_tasks = set()
        self.state.story_import_requests = set()
        self.state.story_shutting_down = False
        self.state.story_seed_factory = lambda: 71
        self.state.story_dlc_version = dlc_provenance(
            self.cfg,
            project_root=self.root,
            waveform_policy=self.harness.controller.waveform_policy,
        )
        self.llm = LocalChapterPlannerClient()
        self.state.llm = self.llm
        self.state.chapter_planner = ChapterPlanner(
            self.llm,
            waveform_registry=self.harness.safety.presets,
            effective_caps={
                channel: self.harness.safety.cap_for(channel)
                for channel in ("A", "B")
            },
            reading_speed_cpm={"slow": 250, "standard": 400, "fast": 600},
            cycle_gap_policy=CycleGapPolicy(),
            safety_adapter=self.harness.safety,
            model_identity=self.llm.model,
            prompt_version="faithful-offline-v1",
            dlc_version=self.state.story_dlc_version,
        )
        self.state.novel_session = NovelSessionController(
            self.harness.controller,
            source_encoding="auto",
            analysis_version="faithful-offline-v1",
            dlc_version=self.state.story_dlc_version,
        )
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", DeprecationWarning)
            with (
                patch.object(main_module, "load_config", return_value=self.cfg),
                patch.object(main_module, "AppState", return_value=self.state),
            ):
                self.app = main_module.make_app()
        self.client = httpx.AsyncClient(
            transport=httpx.ASGITransport(app=self.app, raise_app_exceptions=False),
            base_url="http://testserver",
        )

    async def asyncTearDown(self) -> None:
        await self.client.aclose()
        with suppress(Exception):
            await self.harness.controller.stop()

    async def test_two_chapter_docx_candidate_runs_offline_from_import_to_exact_replay(self):
        """A missing cache, network call, relay send, or replay asset must fail this gate."""
        source_bytes = self._two_chapter_docx()
        source_path = self.root / "fixture.docx"
        source_path.write_bytes(source_bytes)
        story = self.state.story_source_loader.load("fixture.docx", source_bytes)
        candidate_path = self._write_candidate(story)
        logical_actions, logical_clears = self._record_game_loop_boundaries()

        self._run_candidate_cli("validate", source_path, candidate_path)
        self._run_candidate_cli("import", source_path, candidate_path)
        self.assertEqual(self.llm.calls, [])

        imported = await self.client.post(
            "/api/story/import",
            data={"encoding": "auto"},
            files={"file": ("fixture.docx", source_bytes, "application/octet-stream")},
        )
        self.assertEqual(imported.status_code, 200, imported.text)
        source_id = imported.json()["source"]["source_id"]
        self.assertEqual(imported.json()["analysis"]["status"], "ready")

        analysis = await self.client.get(f"/api/story/{source_id}/analysis")
        chapters = await self.client.get(f"/api/story/{source_id}/chapters")
        self.assertEqual(analysis.status_code, 200, analysis.text)
        self.assertEqual(analysis.json()["status"], "ready")
        self.assertEqual(chapters.status_code, 200, chapters.text)
        self.assertEqual(len(chapters.json()["chapters"]), 2)
        self.assertEqual(self.llm.calls, [])

        selected_chapter_id = chapters.json()["chapters"][0]["chapter_id"]
        played = await self.client.post(
            f"/api/story/{source_id}/chapters/{selected_chapter_id}/play",
            json={"speed": "standard"},
        )
        self.assertEqual(played.status_code, 200, played.text)
        self.assertEqual(played.json()["status"], "running")
        self._assert_selected_chapter_request(story)

        await self._advance_until(
            lambda: (
                bool(self.harness.controller.recorded_cycles)
                and self.harness.safety.current["A"] > 0
                and self.harness.safety.current["B"] > 0
            ),
            "the dry-run chapter should execute resolved A/B cycles",
        )
        before_pause = self.state.novel_session.to_state()
        before_strengths = dict(self.harness.safety.current)
        self.assertEqual(before_pause.status.value, "running")
        self.assertTrue(before_pause.current_scene_id)
        self.assertGreater(before_strengths["A"], 0)
        self.assertGreater(before_strengths["B"], 0)
        clear_start = len(logical_clears)
        paused = await self.client.post("/api/story/pause")
        self.assertEqual(paused.json()["status"], "paused")
        self.assertEqual(paused.json()["cursor"], before_pause.cursor)
        self.assertEqual(paused.json()["current_scene_id"], before_pause.current_scene_id)
        pause_clears = logical_clears[clear_start:]
        self.assertEqual([channel for channel, _ in pause_clears], [None])
        pause_executed, pause_dropped = pause_clears[0][1]
        self.assertEqual(pause_dropped, [])
        self.assertEqual(pause_executed[0]["action"], {"op": "stop"})
        self.assertEqual(
            pause_executed[0]["effective"]["channels"],
            {
                "A": {
                    "effective_strength": 0,
                    "pattern": None,
                    "waveform_mode": None,
                },
                "B": {
                    "effective_strength": 0,
                    "pattern": None,
                    "waveform_mode": None,
                },
            },
        )
        self.assertEqual(self.harness.safety.current, {"A": 0, "B": 0})
        resumed = await self.client.post("/api/story/resume", json={"from": "current"})
        self.assertEqual(resumed.json()["status"], "running")
        await self._advance_until(
            lambda: (
                self.state.novel_session.to_state().cursor == before_pause.cursor
                and self.state.novel_session.to_state().current_scene_id
                == before_pause.current_scene_id
            ),
            "resume current should restart the paused safe event",
        )
        self.assertEqual(self.harness.relay.sent_frames, [])

        finished = await self.client.post("/api/story/finish")
        self.assertEqual(finished.status_code, 200, finished.text)
        self.assertEqual(finished.json()["replay"]["status"], "completed")
        replay_id = finished.json()["replay"]["replay_id"]
        archive = self.harness.store.load(replay_id)
        completed_cycles = tuple(cycle for cycle in archive.timeline.cycles if cycle.completed)
        self.assertTrue(completed_cycles)
        expected_scenes = self._expected_scenes(story)
        self.assertEqual(archive.source, source_bytes)
        self.assertEqual(archive.manifest.source_hash, hashlib.sha256(source_bytes).hexdigest())
        self.assertEqual(archive.scenes, expected_scenes)
        self.assertEqual(
            archive.manifest.checksums["scenes.json"], self._json_hash(expected_scenes)
        )
        self.assertEqual(archive.scenes["source_hash"], story.source_sha256)

        replay_action_start = len(logical_actions)
        replayed = await self.client.post(
            f"/api/replays/{replay_id}/play", json={"cursor": 0}
        )
        self.assertEqual(replayed.status_code, 200, replayed.text)
        self.assertEqual(replayed.json()["status"], "replaying")
        await self._advance_until(
            lambda: self.harness.controller.to_state().status is SessionStatus.IDLE,
            "exact replay should complete from the archived timeline",
        )
        self.assertEqual(self.harness.relay.sent_frames, [])
        self.assertEqual([call["schema_name"] for call in self.llm.calls], ["chapter_plan"])
        replay_actions = logical_actions[replay_action_start:]
        self.assertEqual(
            [self._cycle_action_identity(actions) for actions in replay_actions],
            [
                (cycle.channel, cycle.pattern, cycle.requested_strength)
                for cycle in completed_cycles
            ],
        )

    def _two_chapter_docx(self) -> bytes:
        document = Document()
        document.add_heading("第一章", level=1)
        document.add_paragraph("甲乙丙丁戊己庚辛")
        document.add_heading("第二章", level=1)
        document.add_paragraph("壬癸子丑寅卯辰巳")
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()

    def _write_candidate(self, story) -> Path:
        second_chapter_start = story.text.index("第二章")
        bounds = ((0, second_chapter_start), (second_chapter_start, len(story.text)))
        chapters = []
        for index, (start, end) in enumerate(bounds):
            chapters.append(
                {
                    "id": f"ch-{story.source_sha256}-{index + 1:04d}",
                    "index": index,
                    "start_offset": start,
                    "end_offset": end,
                    "title": f"第{index + 1}章",
                    "summary": f"第{index + 1}章的本地摘要。",
                    "scenes": [
                        {
                            "id": f"ch-{story.source_sha256}-{index + 1:04d}-sc-0001",
                            "index": 0,
                            "start_offset": start,
                            "end_offset": end,
                            "summary": f"第{index + 1}章场景。",
                            "pace": 1.0,
                        }
                    ],
                }
            )
        candidate = {
            "source_hash": story.source_sha256,
            "text_length": len(story.text),
            "chapters": chapters,
        }
        directory = self.root / self.cfg["story"]["candidate_dir"]
        directory.mkdir(parents=True)
        path = directory / "mvp2-two-chapter.json"
        path.write_text(json.dumps(candidate, ensure_ascii=False), encoding="utf-8")
        return path

    def _expected_scenes(self, story) -> dict[str, object]:
        # The importer has already proved this cached map against the real source.
        story_map = self.state.story_analysis_store.load(
            offline_analysis_key(story, self.state.story_dlc_version)
        )
        self.assertIsNotNone(story_map)
        return encode_story_map(story_map, schema_version=SCHEMA_VERSION)

    def _assert_selected_chapter_request(self, story) -> None:
        self.assertEqual(len(self.llm.calls), 1)
        call = self.llm.calls[0]
        self.assertEqual(call["schema_name"], "chapter_plan")
        request = json.loads(call["user_content"])
        story_map = self.state.story_analysis_store.load(
            offline_analysis_key(story, self.state.story_dlc_version)
        )
        self.assertIsNotNone(story_map)
        chapter = story_map.chapters[0]
        scene = chapter.scenes[0]
        self.assertEqual(
            request,
            {
                "chapter": {
                    "id": chapter.id,
                    "index": 0,
                    "start_offset": chapter.start_offset,
                    "end_offset": chapter.end_offset,
                    "title": chapter.title,
                    "summary": chapter.summary,
                },
                "chapter_text": story.text[chapter.start_offset : chapter.end_offset],
                "scenes": [
                    {
                        "scene_id": scene.id,
                        "index": 0,
                        "start_offset": scene.start_offset,
                        "end_offset": scene.end_offset,
                        "summary": scene.summary,
                    }
                ],
                "capabilities": {
                    "A": {
                        "modes": ["keep", "set", "stop"],
                        "set_fields": ["base_strength"],
                        "effective_cap": self.harness.safety.cap_for("A"),
                    },
                    "B": {
                        "modes": ["keep", "set", "stop"],
                        "set_fields": ["base_strength"],
                        "effective_cap": self.harness.safety.cap_for("B"),
                    },
                },
                "constraints": {
                    "faithful_mode": True,
                    "preserve_scene_order": True,
                    "one_entry_per_scene": True,
                    "waveform_source": "seeded_resolver",
                },
            },
        )
        second_chapter = story_map.chapters[1]
        forbidden = (
            story.text[second_chapter.start_offset : second_chapter.end_offset],
            second_chapter.id,
            second_chapter.scenes[0].id,
            second_chapter.summary,
            second_chapter.scenes[0].summary,
            "whole_book",
            "other_chapters",
            "chat",
            "camera",
            "microphone",
        )
        for content in (call["system_prompt"], call["user_content"]):
            for forbidden_value in forbidden:
                self.assertNotIn(forbidden_value, content)
        mutated = dict(request)
        mutated["chapter_text"] = request["chapter_text"] + story.text[
            second_chapter.start_offset : second_chapter.end_offset
        ]
        with self.assertRaises(AssertionError):
            self.assertEqual(mutated, request)

    def _record_game_loop_boundaries(self):
        logical_actions: list[tuple[dict[str, object], ...]] = []
        logical_clears: list[tuple[str | None, tuple[list, list]]] = []
        original_timeline_actions = self.harness.loop.execute_timeline_actions
        original_clear_output = self.harness.loop.clear_output

        async def record_timeline_actions(actions, owner_generations):
            logical_actions.append(tuple(deepcopy(actions)))
            return await original_timeline_actions(actions, owner_generations)

        async def record_clear_output(channel=None):
            result = await original_clear_output(channel)
            logical_clears.append((channel, result))
            return result

        self.harness.loop.execute_timeline_actions = record_timeline_actions
        self.harness.loop.clear_output = record_clear_output
        self.addCleanup(
            setattr,
            self.harness.loop,
            "execute_timeline_actions",
            original_timeline_actions,
        )
        self.addCleanup(
            setattr, self.harness.loop, "clear_output", original_clear_output
        )
        return logical_actions, logical_clears

    @staticmethod
    def _cycle_action_identity(actions):
        if len(actions) != 2:
            raise AssertionError("replay must execute a strength and one cycle action")
        strength, cycle = actions
        if strength.get("op") != "hold_strength" or cycle.get("op") != "pulse_cycle":
            raise AssertionError("replay did not execute the recorded cycle actions")
        if strength.get("channel") != cycle.get("channel"):
            raise AssertionError("replay cycle actions used different channels")
        return strength["channel"], cycle.get("pattern"), strength.get("value")

    def _run_candidate_cli(self, action: str, source_path: Path, candidate_path: Path) -> None:
        stdout = io.StringIO()
        stderr = io.StringIO()
        with (
            patch("backend.story.import_analysis.load_config", return_value=self.cfg),
            patch("backend.story.import_analysis.PROJECT_ROOT", self.root),
            redirect_stdout(stdout),
            redirect_stderr(stderr),
        ):
            status = import_analysis_main(
                [action, "--source", str(source_path), "--map", str(candidate_path)]
            )
        self.assertEqual(status, 0, stderr.getvalue())
        self.assertEqual(stderr.getvalue(), "")
        self.assertIn(f"status={'validated' if action == 'validate' else 'imported'}", stdout.getvalue())

    async def _advance_until(self, predicate, message: str) -> None:
        for _ in range(200):
            if predicate():
                return
            remaining = self.harness.clock.next_remaining_ms
            if remaining is not None:
                self.harness.clock.advance(remaining)
            await asyncio.sleep(0)
        self.fail(message)

    @staticmethod
    def _json_hash(value: object) -> str:
        payload = json.dumps(
            value, ensure_ascii=False, sort_keys=True, separators=(",", ":")
        ).encode("utf-8")
        return hashlib.sha256(payload).hexdigest()


if __name__ == "__main__":
    unittest.main()
