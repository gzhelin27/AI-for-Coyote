import hashlib
import io
from dataclasses import FrozenInstanceError
from pathlib import Path
import tempfile
from typing import get_args
import unittest
import zipfile

from docx import Document
import yaml

import backend.story as story_domain
from backend.config import DEFAULTS, load_config
from backend.story.models import ImportedStory
from backend.story.source import StorySourceError, StorySourceLoader


_AMBIGUOUS_GB18030_CASES = (
    ("D2BB", bytes.fromhex("d2bb"), "一"),
    ("D2BBD2B5", bytes.fromhex("d2bbd2b5"), "一业"),
    ("repeated D2BB", bytes.fromhex("d2bb") * 8, "一" * 8),
)
_MULTILINGUAL_TEXTS = ("中文", "😊", "Привет", "café")


class StorySourceTests(unittest.TestCase):
    def test_txt_normalizes_newlines_and_hashes_normalized_text(self):
        original = "甲\r\n乙".encode("utf-8")

        imported = self._load_with_encoding("novel.txt", original, "utf-8")

        self.assertEqual(imported.text, "甲\n乙")
        self.assertEqual(imported.filename, "novel.txt")
        self.assertEqual(imported.extension, ".txt")
        self.assertEqual(
            imported.source_sha256,
            hashlib.sha256("甲\n乙".encode("utf-8")).hexdigest(),
        )
        self.assertEqual(imported.original_bytes, original)
        with self.assertRaises(FrozenInstanceError):
            imported.text = "changed"

    def test_markdown_uses_gb18030_after_utf8_fails_and_strips_text_nuls(self):
        original = "第一章\x00\r第二章".encode("gb18030")

        imported = StorySourceLoader(max_bytes=1024).load("chapter.md", original)

        self.assertEqual(imported.text, "第一章\n第二章")
        self.assertEqual(
            imported.source_sha256,
            hashlib.sha256("第一章\n第二章".encode("utf-8")).hexdigest(),
        )

    def test_equivalent_normalized_text_has_one_identity_across_line_endings_bom_and_encoding(self):
        loader = StorySourceLoader(max_bytes=1024)
        variants = (
            ("lf", "第一章\n正文".encode("utf-8"), "utf-8"),
            ("crlf", "第一章\r\n正文".encode("utf-8"), "utf-8"),
            ("bom", b"\xef\xbb\xbf" + "第一章\n正文".encode("utf-8"), "auto"),
            ("gb18030", "第一章\n正文".encode("gb18030"), "gb18030"),
        )

        imported = [
            loader.load(f"{name}.txt", payload, encoding=encoding)
            for name, payload, encoding in variants
        ]

        self.assertEqual({story.text for story in imported}, {"第一章\n正文"})
        self.assertEqual(len({story.source_sha256 for story in imported}), 1)
        self.assertEqual(imported[0].original_bytes, variants[0][1])
        self.assertNotEqual(imported[0].original_bytes, imported[1].original_bytes)

    def test_docx_container_metadata_does_not_change_normalized_text_identity(self):
        first = io.BytesIO()
        second = io.BytesIO()
        document = Document()
        document.add_paragraph("第一章")
        document.add_paragraph("正文")
        document.save(first)
        document.core_properties.author = "Different container metadata"
        document.save(second)

        loader = StorySourceLoader(max_bytes=1024 * 1024)
        imported_first = loader.load("first.docx", first.getvalue())
        imported_second = loader.load("second.docx", second.getvalue())

        self.assertEqual(imported_first.text, "第一章\n正文")
        self.assertEqual(imported_first.source_sha256, imported_second.source_sha256)
        self.assertNotEqual(imported_first.original_bytes, imported_second.original_bytes)

    def test_story_source_encoding_options_are_public(self):
        encoding_type = getattr(story_domain, "StorySourceEncoding", None)

        self.assertEqual(get_args(encoding_type), ("auto", "utf-8", "gb18030"))

    def test_auto_rejects_ambiguous_gb18030_samples(self):
        for label, payload, _expected in _AMBIGUOUS_GB18030_CASES:
            with self.subTest(label=label), self.assertRaisesRegex(
                StorySourceError, "encoding is ambiguous"
            ):
                self._load_with_encoding("chapter.txt", payload, "auto")

    def test_default_encoding_is_auto(self):
        with self.assertRaisesRegex(StorySourceError, "encoding is ambiguous"):
            StorySourceLoader(max_bytes=1024).load("chapter.txt", bytes.fromhex("d2bb"))

    def test_explicit_gb18030_preserves_ambiguous_chinese_samples(self):
        for label, payload, expected in _AMBIGUOUS_GB18030_CASES:
            with self.subTest(label=label):
                imported = self._load_with_encoding("chapter.txt", payload, "gb18030")

                self.assertEqual(imported.text, expected)

    def test_explicit_encodings_preserve_multilingual_text(self):
        for encoding in ("utf-8", "gb18030"):
            for expected in _MULTILINGUAL_TEXTS:
                with self.subTest(encoding=encoding, expected=expected):
                    imported = self._load_with_encoding(
                        "chapter.txt", expected.encode(encoding), encoding
                    )

                    self.assertEqual(imported.text, expected)

    def test_auto_rejects_bomless_multilingual_utf8_when_decodings_differ(self):
        for text in _MULTILINGUAL_TEXTS:
            with self.subTest(text=text), self.assertRaisesRegex(
                StorySourceError, "encoding is ambiguous"
            ):
                self._load_with_encoding("chapter.txt", text.encode("utf-8"), "auto")

    def test_auto_accepts_a_unique_utf8_decoding(self):
        imported = self._load_with_encoding("chapter.txt", b"\xe0\xa0\x80", "auto")

        self.assertEqual(imported.text, "\u0800")

    def test_utf8_bom_is_definitive_and_stripped_in_auto_and_explicit_utf8(self):
        for expected in ("ASCII chapter", "中", "开篇😊"):
            original = b"\xef\xbb\xbf" + expected.encode("utf-8")
            for encoding in ("auto", "utf-8"):
                with self.subTest(expected=expected, encoding=encoding):
                    imported = self._load_with_encoding("chapter.txt", original, encoding)

                    self.assertEqual(imported.text, expected)

    def test_utf8_bom_rejects_conflicting_explicit_gb18030(self):
        for text in ("ASCII chapter", "中"):
            original = b"\xef\xbb\xbf" + text.encode("utf-8")
            with self.subTest(text=text), self.assertRaisesRegex(
                StorySourceError, "UTF-8 BOM.*conflicts.*gb18030"
            ):
                self._load_with_encoding("chapter.txt", original, "gb18030")

    def test_rejects_invalid_encoding_name(self):
        with self.assertRaisesRegex(StorySourceError, "encoding must be one of"):
            self._load_with_encoding("chapter.txt", b"text", "latin-1")

    def test_wrong_explicit_encoding_fails_without_fallback(self):
        cases = (
            ("utf-8", "第一章".encode("gb18030")),
            ("gb18030", b"\xe0\xa0\x80"),
        )

        for encoding, payload in cases:
            with self.subTest(encoding=encoding), self.assertRaisesRegex(
                StorySourceError, f"not valid {encoding}"
            ):
                self._load_with_encoding("chapter.txt", payload, encoding)

    def test_docx_preserves_heading_order_and_paragraph_text(self):
        document = Document()
        document.add_heading("第一章", level=1)
        document.add_paragraph("雨落在窗前。")
        document.add_heading("第二章", level=1)
        document.add_paragraph("门在黎明前打开。")
        payload = io.BytesIO()
        document.save(payload)
        original = payload.getvalue()

        imported = StorySourceLoader(max_bytes=1024 * 1024).load("novel.docx", original)

        self.assertEqual(imported.text, "第一章\n雨落在窗前。\n第二章\n门在黎明前打开。")
        self.assertEqual(imported.extension, ".docx")
        self.assertEqual(
            imported.source_sha256,
            hashlib.sha256(imported.text.encode("utf-8")).hexdigest(),
        )

    def test_docx_rejects_plain_text_encoding_override(self):
        original = self._generated_docx_bytes()

        for encoding in ("utf-8", "gb18030"):
            with self.subTest(encoding=encoding), self.assertRaisesRegex(
                StorySourceError, "DOCX.*encoding"
            ):
                self._load_with_encoding("novel.docx", original, encoding)

    def test_rejects_unsupported_extension_and_oversize_input(self):
        loader = StorySourceLoader(max_bytes=3)

        with self.assertRaises(StorySourceError):
            loader.load("novel.pdf", b"abc")
        with self.assertRaises(StorySourceError):
            loader.load("novel.txt", b"abcd")

    def test_auto_accepts_identical_ascii_at_exact_byte_limit(self):
        imported = StorySourceLoader(max_bytes=4).load(
            "NOVEL.TXT", b"text", encoding="auto"
        )

        self.assertEqual(imported.filename, "NOVEL.TXT")
        self.assertEqual(imported.extension, ".txt")

    def test_sanitizes_filename_without_using_parent_path(self):
        imported = StorySourceLoader(max_bytes=1024).load(
            str(Path("parent") / ".." / "novel.txt"), b"safe text"
        )

        self.assertEqual(imported.filename, "novel.txt")

    def test_rejects_invalid_filename_bytes_encoding_and_empty_text(self):
        loader = StorySourceLoader(max_bytes=1024)
        invalid_cases = (
            ("", b"text"),
            ("story\x00.txt", b"text"),
            ("story.txt", "not bytes"),
            ("story.txt", b"\xff"),
            ("story.txt", b"\x00\r\n\x00"),
        )

        for filename, payload in invalid_cases:
            with self.subTest(filename=filename, payload=payload), self.assertRaises(StorySourceError):
                loader.load(filename, payload)

    def test_rejects_docx_with_excessive_compressed_expansion_before_extraction(self):
        original = self._generated_docx_bytes()
        payload = io.BytesIO(original)
        with zipfile.ZipFile(payload, "a", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("padding.txt", b"A" * 1_000_000)

        with self.assertRaises(StorySourceError):
            StorySourceLoader(max_bytes=2 * 1024 * 1024).load("novel.docx", payload.getvalue())

    def test_rejects_docx_with_excessive_entry_count_before_extraction(self):
        payload = io.BytesIO(self._generated_docx_bytes())
        with zipfile.ZipFile(payload, "a", compression=zipfile.ZIP_DEFLATED) as archive:
            for number in range(300):
                archive.writestr(f"padding/{number}.txt", b"x")

        with self.assertRaises(StorySourceError):
            StorySourceLoader(max_bytes=2 * 1024 * 1024).load("novel.docx", payload.getvalue())

    def test_rejects_encrypted_docx_before_parser(self):
        encrypted = self._mark_central_directory_encrypted(self._generated_docx_bytes())

        with self.assertRaisesRegex(StorySourceError, "encrypted"):
            StorySourceLoader(max_bytes=1024 * 1024).load("novel.docx", encrypted)

    def test_rejects_docx_with_unsafe_member_name(self):
        payload = io.BytesIO(self._generated_docx_bytes())
        with zipfile.ZipFile(payload, "a", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("../outside.txt", b"not allowed")

        with self.assertRaises(StorySourceError):
            StorySourceLoader(max_bytes=1024 * 1024).load("novel.docx", payload.getvalue())

    def test_rejects_malformed_docx(self):
        with self.assertRaises(StorySourceError):
            StorySourceLoader(max_bytes=1024).load("novel.docx", b"not a zip file")

    @staticmethod
    def _generated_docx_bytes() -> bytes:
        document = Document()
        document.add_paragraph("正常文档")
        payload = io.BytesIO()
        document.save(payload)
        return payload.getvalue()

    @staticmethod
    def _mark_central_directory_encrypted(payload: bytes) -> bytes:
        marked = bytearray(payload)
        position = 0
        while True:
            position = marked.find(b"PK\x01\x02", position)
            if position < 0:
                return bytes(marked)
            flags = int.from_bytes(marked[position + 8:position + 10], "little")
            marked[position + 8:position + 10] = (flags | 0x1).to_bytes(2, "little")
            position += 4

    def _load_with_encoding(
        self, filename: str, payload: bytes, encoding: str
    ) -> ImportedStory:
        try:
            return StorySourceLoader(max_bytes=1024 * 1024).load(
                filename, payload, encoding=encoding
            )
        except TypeError as exc:
            if "encoding" in str(exc):
                self.fail("StorySourceLoader.load does not expose the encoding option")
            raise


class StoryConfigurationTests(unittest.TestCase):
    def test_story_defaults_merge_with_existing_configuration(self):
        with tempfile.TemporaryDirectory() as directory:
            config_path = Path(directory) / "config.yaml"
            config_path.write_text("app:\n  port: 8765\nstory:\n  max_source_mb: 2\n", encoding="utf-8")

            config = load_config(config_path)

        self.assertEqual(config["app"]["port"], 8765)
        self.assertEqual(
            config["story"],
            {
                "import_dir": "data/stories",
                "analysis_dir": "data/story_analysis",
                "candidate_dir": "data/story_candidates",
                "max_source_mb": 2,
                "analysis_prompt_version": "faithful-v1",
                "reading_speed_cpm": {"slow": 250, "standard": 400, "fast": 600},
            },
        )
        self.assertEqual(DEFAULTS["story"]["max_source_mb"], 20)

    def test_story_configuration_normalizes_valid_local_overrides(self):
        config = self._load_config(
            {
                "story": {
                    "import_dir": "data\\stories\\imports",
                    "analysis_dir": "data/analysis/cache",
                    "candidate_dir": "data/story_candidates/custom",
                    "max_source_mb": 2.5,
                    "analysis_prompt_version": " faithful-v2 ",
                    "reading_speed_cpm": {"slow": 250.0, "standard": 401.5, "fast": 600.0},
                }
            }
        )

        self.assertEqual(config["story"]["import_dir"], "data/stories/imports")
        self.assertEqual(config["story"]["analysis_dir"], "data/analysis/cache")
        self.assertEqual(config["story"]["candidate_dir"], "data/story_candidates/custom")
        self.assertEqual(config["story"]["max_source_mb"], 2.5)
        self.assertEqual(config["story"]["analysis_prompt_version"], "faithful-v2")
        self.assertEqual(
            config["story"]["reading_speed_cpm"],
            {"slow": 250, "standard": 401.5, "fast": 600},
        )

    def test_story_configuration_rejects_invalid_values(self):
        invalid_stories = (
            {"max_source_mb": 0},
            {"max_source_mb": True},
            {"max_source_mb": -1},
            {"max_source_mb": float("nan")},
            {"reading_speed_cpm": [250, 400, 600]},
            {"reading_speed_cpm": {"slow": 250, "standard": 400}},
            {"reading_speed_cpm": {"slow": 250, "standard": 400, "fast": 600, "turbo": 800}},
            {"reading_speed_cpm": {"slow": 0, "standard": 400, "fast": 600}},
            {"reading_speed_cpm": {"slow": True, "standard": 400, "fast": 600}},
            {"reading_speed_cpm": {"slow": 250, "standard": float("inf"), "fast": 600}},
            {"import_dir": "../stories"},
            {"analysis_dir": "C:\\story-analysis"},
            {"candidate_dir": "../story-candidates"},
            {"import_dir": "stories/imports"},
            {"analysis_dir": "data"},
            {"import_dir": ""},
            {"analysis_prompt_version": "   "},
            {"unknown": "value"},
            {"reading_speed_cpm": {"slow": 400, "standard": 250, "fast": 600}},
            {"reading_speed_cpm": {"slow": 250, "standard": 600, "fast": 600}},
        )

        for story in invalid_stories:
            with self.subTest(story=story), self.assertRaises(ValueError):
                self._load_config({"story": story})

    @staticmethod
    def _load_config(overrides: dict) -> dict:
        with tempfile.TemporaryDirectory() as directory:
            config_path = Path(directory) / "config.yaml"
            config_path.write_text(yaml.safe_dump(overrides), encoding="utf-8")
            return load_config(config_path)


if __name__ == "__main__":
    unittest.main()
